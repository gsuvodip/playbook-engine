from openpyxl import load_workbook

def process_playbook(playbook_path):
    print(f"Processing playbook: {playbook_path}")

def load_playbook(playbook_path):
    wb = load_workbook(playbook_path, data_only=True)
    sheet_names = wb.sheetnames
    # print(f"Sheets: {sheet_names}")
    return sheet_names